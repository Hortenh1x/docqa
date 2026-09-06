"""Build upload-ready corpus files from corpus/*.md sources.

Most documents become PDFs (exercises the page-aware parser and looks like a real
archive); two stay Markdown and two become DOCX — the multi-format demo. Front matter
is rendered as a visible header so version/effective-date facts survive into chunks
(they matter for the version-conflict trap).

PDF rendering uses PyMuPDF's Story (no pandoc/weasyprint needed). Markdown tables are
wrapped as preformatted blocks so the pipe layout survives PDF text extraction and the
chunker keeps tables atomic.

Usage: uv run python -m scripts.build_corpus
       uv run python -m scripts.build_corpus --src corpus/large/docs --out corpus/large/build \
           --manifest corpus/large/manifest.yaml
"""

import argparse
import re
import shutil
from pathlib import Path

import fitz
import markdown as md_lib
import yaml
from docx import Document as DocxBuilder

CORPUS = Path("corpus")
OUT = CORPUS / "build"

KEEP_MD = {"POL-009", "OPS-001"}
AS_DOCX = {"HR-002", "POL-007"}

_CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.45; color: #1a1a1a; }
h1 { font-size: 20pt; margin: 0 0 6pt 0; }
h2 { font-size: 15pt; margin: 14pt 0 4pt 0; }
h3 { font-size: 13.5pt; margin: 10pt 0 3pt 0; }
p { margin: 4pt 0; }
ul { margin: 4pt 0 4pt 14pt; }
pre { font-family: monospace; font-size: 9pt; line-height: 1.35; }
.meta { color: #555555; font-size: 9pt; }
"""


def split_front_matter(source: str) -> tuple[dict[str, str], str]:
    match = re.match(r"^---\n(.*?)\n---\n", source, re.DOTALL)
    if not match:
        return {}, source
    meta: dict[str, str] = {}
    for line in match.group(1).splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            meta[key.strip()] = value.strip().strip('"')
    return meta, source[match.end() :]


def meta_header_lines(meta: dict[str, str]) -> list[str]:
    parts = []
    for key in ("doc_id", "version", "effective_date", "supersedes", "owner", "classification"):
        if key in meta:
            label = key.replace("_", " ").title()
            parts.append(f"{label}: {meta[key]}")
    return parts


def _fence_tables(body: str) -> str:
    """Wrap markdown tables in code fences so pipes survive PDF text extraction."""
    lines = body.splitlines()
    out: list[str] = []
    in_table = False
    for line in lines:
        is_row = line.lstrip().startswith("|")
        if is_row and not in_table:
            out.append("```")
            in_table = True
        if not is_row and in_table:
            out.append("```")
            in_table = False
        out.append(line)
    if in_table:
        out.append("```")
    return "\n".join(out)


def access_line(meta: dict[str, str]) -> str | None:
    """Whole-document restriction from front matter, rendered as its own paragraph so the
    ingestion parser sees a standalone "Access: … only" marker before the first heading."""
    value = meta.get("access")
    return f"Access: {value}" if value else None


def build_pdf(meta: dict[str, str], body: str, dest: Path) -> None:
    html_body = md_lib.markdown(_fence_tables(body), extensions=["fenced_code"])
    header = " · ".join(meta_header_lines(meta))
    html = f'<p class="meta">{header}</p>\n'
    if access := access_line(meta):
        html += f'<p class="meta">{access}</p>\n'
    html += html_body

    story = fitz.Story(html=html, user_css=_CSS)
    writer = fitz.DocumentWriter(str(dest))
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (54, 56, -54, -56)
    more = 1
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()


def build_docx(meta: dict[str, str], body: str, dest: Path) -> None:
    doc = DocxBuilder()
    for line in meta_header_lines(meta):
        doc.add_paragraph(line)
    if access := access_line(meta):
        doc.add_paragraph(access)

    table_rows: list[list[str]] = []

    def flush_table() -> None:
        nonlocal table_rows
        rows = [r for r in table_rows if not set("".join(r)) <= {"-", " ", ":"}]
        if rows:
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    table.rows[i].cells[j].text = cell
        table_rows = []

    for block in body.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.lstrip().startswith("|"):
            for line in block.splitlines():
                table_rows.append([c.strip() for c in line.strip().strip("|").split("|")])
            flush_table()
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", block.splitlines()[0])
        if heading:
            doc.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 4))
            rest = "\n".join(block.splitlines()[1:]).strip()
            if rest:
                doc.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", rest))
            continue
        doc.add_paragraph(re.sub(r"\*\*(.+?)\*\*", r"\1", block))
    doc.save(str(dest))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--src", type=Path, default=CORPUS, help="directory with *.md sources")
    parser.add_argument(
        "--out", type=Path, default=None, help="output directory (default: <src>/build)"
    )
    parser.add_argument(
        "--manifest",
        type=Path,
        default=None,
        help="corpus v2 manifest.yaml: takes the output format per document from it",
    )
    args = parser.parse_args()
    src: Path = args.src
    out: Path = args.out or (src / "build")
    formats: dict[str, str] = {}
    if args.manifest:
        for doc in yaml.safe_load(args.manifest.read_text(encoding="utf-8")):
            formats[doc["doc_id"]] = doc.get("format", "pdf")

    out.mkdir(parents=True, exist_ok=True)
    built = []
    for source_path in sorted(src.glob("*.md")):
        if source_path.name in ("FACTS.md", "style.md"):
            continue
        doc_id = source_path.stem
        meta, body = split_front_matter(source_path.read_text(encoding="utf-8"))
        if formats:
            fmt = formats.get(doc_id, "pdf")
        else:
            fmt = "md" if doc_id in KEEP_MD else "docx" if doc_id in AS_DOCX else "pdf"
        if fmt == "md":
            dest = out / f"{doc_id}.md"
            shutil.copyfile(source_path, dest)
        elif fmt == "docx":
            dest = out / f"{doc_id}.docx"
            build_docx(meta, body, dest)
        else:
            dest = out / f"{doc_id}.pdf"
            build_pdf(meta, body, dest)
        built.append(dest.name)
    print(f"built {len(built)} files into {out}/")
    if len(built) <= 40:
        for name in built:
            print(f"  {name}")


if __name__ == "__main__":
    main()
